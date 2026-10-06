"""src/kp_export.py — экспорт панели КП в данные для Excel.

build_kp_sheets(kp_items, date_str) возвращает {"КП": [[...], ...],
"Служебный": [[...], ...]} — построчные данные обоих листов, без
оформления (headless, без openpyxl/tkinter, тестируется напрямую);
write_xlsx(sheets, path) — единственное место оформления книги (openpyxl
импортируется внутри функции), gui.py только вызывает её.

Лист «КП» повторяет структуру бизнес-шаблона коммерческого предложения
(3 строки шапки — заголовок+дата, примечание про НДС, пустая строка —
строка колонок (12 штук), построчные позиции, пустая строка, «ИТОГО без
НДС:»/«ВСЕГО с НДС:»); ставка НДС — core.VAT_RATE.

Наименование/Давление: для точной позиции по маркировке заявки (маршрут
resolve) — маркировка клиента (без количества, стандарта и стали) +
« (сталь)», давление — из заявки; для остальных — маркировка прайса с
подставленным давлением заявки вместо «(Ру)», если оно попадает в диапазон
строки, иначе как в прайсе. Маркировка прайса и диапазон каталога — на
листе «Служебный».
Лист «Служебный» — внутренние данные для разбора КП менеджером: SKU,
позиция заявки, примечание-аналог, источник/координата происхождения строки
каталога, «Ограничитель» (matching.restrictor_kind), «Маркировка прайса»,
«Диапазон давления каталога».
"""
import re

import core
import matching

VAT_RATE = core.VAT_RATE  # единственный источник — core.VAT_RATE (итерация 47)

KP_HEADERS = ["№", "Наименование", "Стандарт", "DN/Размер", "Давление",
              "Толщина", "Материал", "Кол-во", "Цена без НДС",
              "Сумма без НДС", f"НДС {int(VAT_RATE * 100)}%", "Сумма с НДС"]

SERVICE_HEADERS = ["SKU", "Поз. заявки", "Примечание (аналог)", "Источник",
                    "Координата", "Ограничитель", "Маркировка прайса",
                    "Диапазон давления каталога"]

_NUM_TXT_RE = re.compile(r"-?\d+(?:[.,]\d+)?")
_STD_WORD_RE = re.compile(r"(?<!\w)(?:ГОСТ|ОСТ)(?:\s+Р)?(?:\s*\d[\d.\-]*)?(?!\w)", re.IGNORECASE)
_ASME_WORD_RE = re.compile(r"(?<!\w)(?:ASME|ANSI)(?:\s*B\s*16[.,]\d+)?(?!\w)", re.IGNORECASE)


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and v != v) or str(v).strip() in ("", "nan", "None")


def as_number(v):
    """Число, если значение числовое («250», "2.5", 16.0 -> 250 / 2.5 / 16),
    иначе исходное значение (диапазон «10.0 - 160.0», дюймы — текст)."""
    if _blank(v):
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        f = float(v)
    else:
        t = str(v).strip()
        if not _NUM_TXT_RE.fullmatch(t):
            return v
        f = float(t.replace(",", "."))
    return int(f) if f.is_integer() else f


def _fmt_num(x) -> str:
    return f"{x:g}"


def client_marking(raw) -> str:
    """Маркировка из строки заявки клиента: без количества/единиц
    (matching.split_qty), стандарта («ГОСТ Р 52376-2005», «ОСТ 26.260.454-99»,
    «ASME B16.20») и кода стали (сталь идёт в скобках из строки каталога)."""
    s, _qty = matching.split_qty(raw)
    s = matching.STD_NUM_RE.sub(" ", s)
    s = _STD_WORD_RE.sub(" ", s)
    s = _ASME_WORD_RE.sub(" ", s)
    s = matching.STEEL_RE.sub(" ", s)
    s = re.sub(r"\(\s*\)", " ", s)
    return re.sub(r"\s+", " ", s).strip(" -,;")


def _mpa_units(it) -> bool:
    return (str(it.get("category") or "") in core.PRESSURE_TEXT_MPA_CATEGORIES
            or str(it.get("standard") or "") == "ОСТ")


def request_pressure(it):
    """Давление заявки в единицах строки каталога (кгс; МПа у ОСТ и колец;
    класс у ASME) — число или текст «lo-hi»; None, если в заявке давления нет
    или оно вне диапазона строки (аналог другого давления)."""
    req = it.get("request") or {}
    lo = hi = None
    if req.get("pn") is not None:
        lo = hi = float(req["pn"])
    elif req.get("pn_lo") is not None and req.get("pn_hi") is not None:
        lo, hi = float(req["pn_lo"]), float(req["pn_hi"])
    elif req.get("class_lo") is not None:
        c_lo, c_hi = req["class_lo"], req.get("class_hi") or req["class_lo"]
        k_lo, k_hi = core.ANSI_CLASS_TO_KGS.get(c_lo), core.ANSI_CLASS_TO_KGS.get(c_hi)
        if k_lo is None or k_hi is None or not _in_range(it, k_lo, k_hi):
            return None
        return f"CL{c_lo}" if c_lo == c_hi else f"CL{c_lo}-{c_hi}"
    if lo is None or not _in_range(it, lo, hi):
        return None
    if _mpa_units(it):
        lo, hi = lo / 10.0, hi / 10.0
    return as_number(lo) if lo == hi else f"{_fmt_num(lo)}-{_fmt_num(hi)}"


def _in_range(it, lo, hi) -> bool:
    c_lo, c_hi = it.get("catalog_kgs_lo"), it.get("catalog_kgs_hi")
    if c_lo is None and c_hi is None:
        return True
    c_lo = c_lo if c_lo is not None else c_hi
    c_hi = c_hi if c_hi is not None else c_lo
    return c_lo - 1e-9 <= lo and hi <= c_hi + 1e-9


_MAT_PAREN_RE = re.compile(r"\(([^()]+)\)\s*$")


def kp_material(material) -> str:
    """Сталь для наименования КП (итерация 50b): «SS321 (08Х18Н10Т)» ->
    «08Х18Н10Т» (марка ГОСТ в скобках заменяет обозначение целиком — без
    двойных скобок), «SS304» -> «SS304», пусто -> «»."""
    if _blank(material):
        return ""
    mat = str(material).strip()
    m = _MAT_PAREN_RE.search(mat)
    return m.group(1).strip() if m else mat


def kp_name(it) -> str:
    mark = "" if _blank(it.get("mark")) else str(it.get("mark"))
    req = it.get("request") or {}
    base = ""
    if req.get("route") == "resolve" and not it.get("is_analog") and req.get("raw"):
        base = client_marking(req["raw"])
        if base and core.GROUP_FULLMETAL in mark and "НАВИТ" not in base.upper():
            base = f"{base} {core.GROUP_FULLMETAL}"   # два продукта — не терять различие
    if not base:
        base = mark
        p = request_pressure(it)
        if p is not None and matching.PN_PLACEHOLDER_CANON in base:
            txt = _fmt_num(p) if isinstance(p, (int, float)) else str(p)
            base = base.replace(matching.PN_PLACEHOLDER_CANON, txt.replace(".", ","))
    mat = kp_material(it.get("material"))
    return f"{base} ({mat})" if mat else base


def kp_pressure(it):
    p = request_pressure(it)
    return p if p is not None else as_number(it.get("pn"))


def build_kp_sheets(kp_items, date_str) -> dict:
    kp_rows = [
        ["КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ"] + [None] * 10 + [date_str],
        [f"Цены без НДС. НДС {int(VAT_RATE * 100)}% добавляется в итог."] + [None] * 11,
        [None] * 12,
        list(KP_HEADERS),
    ]

    total_wo = 0.0
    for i, it in enumerate(kp_items, start=1):
        qty = it.get("qty") or 1
        price = it.get("price") or 0.0
        s_wo = qty * price
        vat = s_wo * VAT_RATE
        s_with = s_wo * (1 + VAT_RATE)
        total_wo += s_wo
        kp_rows.append([i, kp_name(it), it.get("standard"), as_number(it.get("dn")),
                         kp_pressure(it), as_number(it.get("thickness")),
                         it.get("material"), as_number(qty),
                         round(price, 2), round(s_wo, 2), round(vat, 2),
                         round(s_with, 2)])

    total_vat = total_wo * VAT_RATE
    total_with = total_wo * (1 + VAT_RATE)
    kp_rows.append([None] * 12)
    kp_rows.append([None, None, None, None, None, None, "ИТОГО без НДС:", None,
                     round(total_wo, 2), round(total_wo, 2), round(total_vat, 2),
                     round(total_with, 2)])
    kp_rows.append([None, None, None, None, None, None, "ВСЕГО с НДС:", None,
                     None, None, None, round(total_with, 2)])

    svc_rows = [list(SERVICE_HEADERS)]
    for it in kp_items:
        restr = matching.restrictor_kind({
            "Маркировка_в_прайсе": it.get("mark"),
            "Стандарт": it.get("standard"),
            "Тип_изделия": it.get("product_type"),
        })
        svc_rows.append([it.get("sku"), it.get("line_no", "—"),
                          it.get("reason", ""), it.get("source"),
                          it.get("coord"), restr, it.get("mark"), it.get("pn")])

    return {"КП": kp_rows, "Служебный": svc_rows}


def write_xlsx(sheets, out_path):
    """Книга КП с оформлением (итерация 50): числа — числами (build_kp_sheets
    уже отдаёт int/float), заголовок в A1 объединён A1:K1 (дата — L1),
    примечание НДС — A2:L2, ширина колонок — по строке колонок и позициям
    (шапка её не растягивает: колонка A — по «№»)."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "КП"
    for row in sheets["КП"]:
        ws.append(row)
    ws.merge_cells("A1:K1")
    ws.merge_cells("A2:L2")
    ws["A1"].font = Font(bold=True, size=14)
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = header_fill
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for c in ws[ws.max_row - 1]:
        c.font = Font(bold=True)
    _autowidth(ws, first_row=4)

    ws2 = wb.create_sheet("Служебный")
    for row in sheets["Служебный"]:
        ws2.append(row)
    for c in ws2[1]:
        c.font = Font(bold=True)
    _autowidth(ws2, first_row=1)
    wb.save(out_path)
    return out_path


def _autowidth(ws, first_row):
    for col in ws.iter_cols(min_row=first_row):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 4), 40)
